import os
import re
import json
import shutil
from pathlib import Path
from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import openpyxl
from docx.opc.constants import RELATIONSHIP_TYPE as RT

# ==================== 配置 ====================
PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = os.getenv("RAG_DATA_DIR", str(PROJECT_ROOT / "data"))
PRE_DATA_DIR = Path(os.getenv("RAG_PRE_DATA_DIR", str(PROJECT_ROOT / "pre_data" / "product")))

MD_DIRS = [
    {"path": os.getenv("RAG_EXTRACT_INPUT", str(Path(DATA_DIR) / "product")), "with_images": False},
]

OUTPUT_JSONL = os.getenv("RAG_EXTRACT_OUTPUT", str(PRE_DATA_DIR / "all_texts.jsonl"))
IMG_DIR = os.getenv("RAG_IMAGE_DIR", str(PRE_DATA_DIR / "img"))

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}

# ==================== 图片过滤 ====================

def is_useful_image(img_path: str) -> bool:
    p = Path(img_path)
    if not p.exists():
        return False
    try:
        if p.stat().st_size / 1024 < 15:
            return False
    except:
        return False
    try:
        with Image.open(p) as im:
            w, h = im.size
            if w < 200 or h < 200:
                return False
            if im.mode in ("RGBA", "LA"):
                alpha = im.getchannel("A")
                transparent_pixels = sum(1 for pixel in alpha.getdata() if pixel < 10)
                if transparent_pixels / (w * h) > 0.5:
                    return False
    except:
        return False
    return True

# ==================== 工具函数 ====================

def ensure_dirs():
    with open(OUTPUT_JSONL, "w", encoding="utf-8") as f:
        pass
    shutil.rmtree(IMG_DIR, ignore_errors=True)
    os.makedirs(IMG_DIR, exist_ok=True)

def write_record(path: str, ext: str, text: str, image_paths):
    rec = {
        "file_path": path,
        "ext": ext,
        "text": text,
        "image_paths": image_paths or []
    }
    with open(OUTPUT_JSONL, "a", encoding="utf-8") as out_f:
        out_f.write(json.dumps(rec, ensure_ascii=False) + "\n")

def safe_copy(src_path: str, dest_dir: str, tag_prefix: str, idx: int) -> str:
    ext = os.path.splitext(src_path)[1].lower()
    tag = tag_prefix.replace(" ", "").replace("/", "")
    dest_name = f"{tag}_{idx:03d}{ext}"
    dest_path = os.path.join(dest_dir, dest_name)
    shutil.copy2(src_path, dest_path)
    return dest_path

def dedup_keep_order(seq):
    seen = set()
    out = []
    for x in seq:
        if x not in seen:
            seen.add(x)
            out.append(x)
    return out

# ==================== 文本提取 ====================

def extract_text_from_docx(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join((cell.text or "").strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts)

def extract_text_from_pptx(path: str) -> str:
    prs = Presentation(path)
    parts = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            parts.append(f"--- Slide {slide_idx} ---")
            parts.append("\n".join(slide_text))
    return "\n".join(parts)

def extract_text_from_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            row_vals = [str(cell) for cell in row if cell is not None]
            if row_vals:
                parts.append("\t".join(row_vals))
    return "\n".join(parts)

# ==================== 图片提取 ====================

def extract_images_from_docx(path: str):
    doc = Document(path)
    img_files = []
    idx = 1
    for rel in doc.part.rels.values():
        if rel.reltype == RT.IMAGE:
            try:
                blob = rel.target_part.blob
                tag = os.path.splitext(os.path.basename(path))[0]
                ext = os.path.splitext(rel.target_part.partname.filename)[1] \
                    if hasattr(rel.target_part, "partname") else ".png"
                if not ext or len(ext) > 6: ext = ".png"
                dest_name = f"{tag}_{idx:03d}{ext}"
                dest_path = os.path.join(IMG_DIR, dest_name)
                with open(dest_path, "wb") as f:
                    f.write(blob)
                img_files.append(dest_path)
                idx += 1
            except Exception as e:
                print(f"ERROR extracting image from DOCX {path}: {e}")
    return img_files

def extract_images_from_pptx(path: str):
    prs = Presentation(path)
    img_files = []
    idx = 1
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = shape.image
                    tag = os.path.splitext(os.path.basename(path))[0]
                    ext = os.path.splitext(img.filename)[1] if hasattr(img, "filename") else ".png"
                    if not ext or len(ext) > 6: ext = ".png"
                    dest_name = f"{tag}_{idx:03d}{ext}"
                    dest_path = os.path.join(IMG_DIR, dest_name)
                    with open(dest_path, "wb") as f:
                        f.write(img.blob)
                    img_files.append(dest_path)
                    idx += 1
                except Exception as e:
                    print(f"ERROR extracting image from PPTX {path}: {e}")
    return img_files

# ==================== 修改的 MD 段落级提取 ====================

def process_md_dir_paragraph_level(md_dir: str, with_images: bool):
    for root, dirs, files in os.walk(md_dir):
        for name in files:
            if os.path.splitext(name)[1].lower() != ".md":
                continue
            path = os.path.join(root, name)
            print(f"Extracting md file (paragraph-level): {path}")
            try:
                with open(path, "r", encoding="utf-8") as f:
                    content = f.read()

                # 按空行拆段
                paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]

                para_idx = 1
                for para in paragraphs:
                    image_paths = []
                    if with_images:
                        # 找该段落的图片引用
                        for p in re.findall(r'(?:!\[.*?\]\(|<img.*?src=")(.*?)(?:"|\))', para):
                            if re.match(r'^(https?:)?//', p.strip()):
                                continue
                            if not os.path.isabs(p):
                                p = os.path.normpath(os.path.join(os.path.dirname(path), p))
                            if os.path.exists(p) and os.path.splitext(p)[1].lower() in IMAGE_EXTS:
                                copied = safe_copy(p, IMG_DIR, f"{os.path.basename(path)}_p{para_idx}", len(image_paths)+1)
                                if is_useful_image(copied):
                                    image_paths.append(copied)

                    if para.strip() or image_paths:
                        write_record(path, ".md", para, image_paths)
                    para_idx += 1

            except Exception as e:
                print(f"ERROR processing {path}: {e}")

# ==================== Office 文件处理 ====================

def process_office_files(root_dir: str):
    for root, dirs, files in os.walk(root_dir):
        for name in files:
            ext = os.path.splitext(name)[1].lower()
            if ext not in {".docx", ".pptx", ".xlsx"}:
                continue
            path = os.path.join(root, name)
            print(f"Extracting office file: {path}")
            try:
                if ext == ".docx":
                    text = extract_text_from_docx(path)
                    images = [p for p in extract_images_from_docx(path) if is_useful_image(p)]
                elif ext == ".pptx":
                    text = extract_text_from_pptx(path)
                    images = [p for p in extract_images_from_pptx(path) if is_useful_image(p)]
                elif ext == ".xlsx":
                    text = extract_text_from_xlsx(path)
                    images = []
                else:
                    continue

                if text.strip() or images:
                    write_record(path, ext, text, images)
                else:
                    print(f"WARNING: empty text and no image from {path}")
            except Exception as e:
                print(f"ERROR processing {path}: {e}")

# ==================== 主入口 ====================

def main():
    ensure_dirs()

    # Office 文件
    process_office_files(DATA_DIR)

    # 多个 MD 目录（段落级处理）
    for item in MD_DIRS:
        p = item["path"]
        w = bool(item.get("with_images", False))
        if os.path.isdir(p):
            process_md_dir_paragraph_level(p, with_images=w)
        else:
            print(f"WARNING: md dir not found: {p}")

    print(f"✅ 段落级文本与图片抽取完成：{OUTPUT_JSONL}")
    print(f"📂 图片保存到：{IMG_DIR}")

if __name__ == "__main__":
    main()
